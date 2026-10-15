"""Refresh the files that ship in `src/dagobert/data/`, for maintainers before a release.

    uv run python scripts/refresh_shipped.py

Writes the series of the Bundesbank and the ECB as delivered, the extracts of the West German
consumer prices and of the gold price of the World Bank, and the newest Destatis life table, and
prints what changed. Then update the dates in `src/dagobert/data/SOURCES.md`. openpyxl comes with
the dev group, the package does not need it.
"""

from __future__ import annotations

import csv
import io
import re
import time
from datetime import date, datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import TYPE_CHECKING

from babel.dates import get_month_names
from openpyxl import load_workbook

from dagobert import series
from dagobert.files import atomic_write

if TYPE_CHECKING:
    from collections.abc import Iterable

    from openpyxl.worksheet.worksheet import Worksheet

DESTATIS = "https://www.destatis.de"
LIFE_TABLES = (
    f"{DESTATIS}/DE/Themen/Gesellschaft-Umwelt/Bevoelkerung/Sterbefaelle-Lebenserwartung/"
    "_inhalt.html"
)
LIFE_TABLE_LINK = re.compile(
    r'<a href="([^"]*?statistischer-bericht-sterbetafeln[^"]*?\.xlsx[^"]*)"[^>]*>'
    r"\s*Statistischer Bericht - Sterbetafeln - (\d{4}/\d{4})"
)
# "Preisindizes für die Lebenshaltung - Lange Reihen von 1948 bis 1999", final since 2025-08
WEST = (
    f"{DESTATIS}/DE/Themen/Wirtschaft/Preise/Verbraucherpreisindex/Publikationen/"
    "Downloads-Verbraucherpreise/"
    "statistischer-bericht-preisindizes-lebenshaltung-lange-reihen-5611111999005.xlsx"
    "?__blob=publicationFile&v=5"
)
MONTHS = {name: month for month, name in get_month_names("wide", locale="de").items()}
PAUSE = 5  # seconds between two requests to a web service: the Bundesbank refuses a burst
PINK_SHEET = re.compile(r'href="([^"]*/CMO-Historical-Data-Monthly\.xlsx)"')
UPDATED = re.compile(r"Updated on (\w+ \d{2}, \d{4})")
STALE = 62  # days: the World Bank updates the workbook early each month


def get(url: str) -> bytes:
    r = series.http_get(url)
    if r.status != 200:
        msg = f"HTTP {r.status} ({r.type}) from {url}"
        raise series.FetchError(msg)
    return r.body


def write(path: Path, text: str) -> None:
    """Write and print how many lines changed: a changed line once, a moved or repeated one too."""
    old = path.read_text(encoding="utf-8").splitlines() if path.exists() else []
    atomic_write(path, text)
    changed = sum(
        max(i2 - i1, j2 - j1)
        for tag, i1, i2, j1, j2 in SequenceMatcher(
            None, old, text.splitlines(), autojunk=False
        ).get_opcodes()
        if tag != "equal"
    )
    print(f"{path.name}: {f'{changed} lines changed' if changed else 'unchanged'}")


def rows(sheet: Worksheet) -> list[dict[str, object]]:
    """The rows of a `csv-` sheet by the names of its header, stripped."""
    cells = sheet.iter_rows(values_only=True)
    header = [str(name).strip() for name in next(cells)]
    return [dict(zip(header, row, strict=True)) for row in cells]


def csv_text(header: tuple[str, ...], lines: Iterable[tuple[object, ...]]) -> str:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(header)
    writer.writerows(lines)
    return out.getvalue()


def target(s: series.Series) -> Path:
    if s.shipped is None:
        msg = f"{s.name} ships no file"
        raise ValueError(msg)
    return series.SHIPPED / s.shipped


def delivered() -> None:
    """The series that ship as their web service delivered them, whole and byte for byte: the
    terms of the Bundesbank and the reuse policy of the ECB forbid a modification."""
    sources = [(s, d) for s in series.SERIES.values() if (d := s.delivered)]
    for i, (s, source) in enumerate(sources):
        time.sleep(PAUSE if i else 0)
        text = source.raw(series.http_get).decode("utf-8")  # keeps the byte order mark
        if not s.parse(text):
            msg = f"{s.name}: no rows in the answer of {source.agency}"
            raise series.FetchError(msg)
        write(target(s), text)


def west() -> None:
    """The index of all private households of the former West Germany, as published (one
    decimal), from its first month 1962-01."""
    sheet = load_workbook(io.BytesIO(get(WEST)), read_only=True)["csv-611xx-b03"]
    points = sorted(
        (int(str(r["Jahr"])), MONTHS[str(r["Monat"])], r["Wert"])
        for r in rows(sheet)
        if r["Statistik"] == "Preisindex für die Lebenshaltung"
        and r["Gebiet"] == "Früheres Bundesgebiet"
        and r["Merkmal_1"] == "Alle privaten Haushalte"
        and r["Wert"] != "-"
    )
    lines = []
    for year, month, value in points:
        if not isinstance(value, int | float) or round(value, 1) != value:
            msg = f"{year}-{month:02d}: {value!r} is no index with one decimal"
            raise ValueError(msg)
        lines.append((f"{year}-{month:02d}", f"{value:.1f}"))
    write(target(series.SERIES["vpi-west"]), csv_text(("month", "value"), lines))


def gold() -> None:
    """The newest monthly workbook of the Pink Sheet, linked from the commodity markets page of
    the World Bank. Names the sources of the gold prices that the source note of `series` must
    list when they changed: the terms of the World Bank ask for them."""
    links = set(PINK_SHEET.findall(get(series.WORLD_BANK).decode("utf-8")))
    if len(links) != 1:
        msg = f"{len(links)} links to the monthly prices on {series.WORLD_BANK}, not one"
        raise ValueError(msg)
    book = load_workbook(io.BytesIO(get(links.pop())), read_only=True)
    entry = series.SERIES["gold"]
    write(target(entry), gold_csv(book["Monthly Prices"]))
    cells = ([str(c) for c in row if c is not None] for row in book["Description"].values)
    sources = next((c[2] for c in cells if len(c) > 2 and c[1].startswith("Gold")), None)
    if sources is None:
        msg = "no Gold row in sheet Description of the World Bank workbook"
        raise ValueError(msg)
    if sources not in entry.attribution:
        print(f"Set the sources in the source note of gold (series.py, SOURCES.md) to {sources}")


def gold_csv(prices: Worksheet) -> str:
    """The column Gold of the sheet Monthly Prices, US dollars per troy ounce from 1960-01, values
    as stored. Refuses a workbook more than `STALE` days old: the link of the page then points to
    an old document."""
    sheet = list(prices.values)
    updated = next((m[1] for row in sheet[:4] if (m := UPDATED.fullmatch(str(row[0])))), "")
    if not updated or (date.today() - datetime.strptime(updated, "%B %d, %Y").date()).days > STALE:
        msg = f"the monthly prices were updated on {updated or 'an unknown day'}: an old workbook"
        raise ValueError(msg)
    column = sheet[4].index("Gold")
    if sheet[5][column] != "($/troy oz)":
        msg = f"the column Gold is in {sheet[5][column]}, not in $/troy oz"
        raise ValueError(msg)
    lines = []
    for row in sheet[6:]:
        month, value = str(row[0]), row[column]
        if not re.fullmatch(r"\d{4}M\d{2}", month) or not isinstance(value, int | float):
            msg = f"unexpected row of the monthly prices: {month} {value!r}"
            raise ValueError(msg)
        lines.append((month.replace("M", "-"), value))
    text = csv_text(("month", "value"), lines)
    if found := series.gap(series.parse_month_value(text)):
        msg = f"gold: {found}"
        raise ValueError(msg)
    return text


def qx(sheet: Worksheet, period: str) -> dict[int, float]:
    out = {}
    for r in rows(sheet):
        age, value = r["Alter"], r["qx"]
        if r["Jahre"] != period or not isinstance(age, int) or not isinstance(value, float):
            msg = f"unexpected row of the life table {period}: {r}"
            raise ValueError(msg)
        out[age] = value
    return out


def copyright_line(sheet: Worksheet) -> str:
    """The source note of a Destatis report, `© Statistisches Bundesamt (Destatis), <year>`."""
    cells = (str(c) for row in sheet.iter_rows(values_only=True) for c in row if c is not None)
    if line := next((c for c in cells if c.startswith("©")), None):
        return line
    msg = "no © line in the sheet of the Destatis report"
    raise ValueError(msg)


def life_table() -> None:
    """The newest period life table for Germany: qx by age, male and female. Names the report
    and the source note that `series` must give when they changed."""
    links = LIFE_TABLE_LINK.findall(get(LIFE_TABLES).decode("utf-8"))
    if not links:
        msg = f"no link to a Statistischer Bericht - Sterbetafeln on {LIFE_TABLES}"
        raise ValueError(msg)
    href, period = max(links, key=lambda link: link[1])
    href = href.replace("&amp;", "&")
    book = load_workbook(io.BytesIO(get(DESTATIS + href)), read_only=True)
    male, female = (qx(book[f"csv-12613-b0{n}"], period) for n in (1, 2))
    if sorted(male) != list(range(len(male))) or sorted(female) != sorted(male):
        msg = f"the life table {period} has gaps in its ages"
        raise ValueError(msg)
    path = series.SHIPPED / f"sterbetafel-{period.replace('/', '-')}.csv"
    lines = [(age, repr(male[age]), repr(female[age])) for age in sorted(male)]
    write(path, csv_text(("age", "qx_m", "qx_f"), lines))
    for old in series.SHIPPED.glob("sterbetafel-*.csv"):
        if old != path:
            old.unlink()
            print(f"{old.name}: removed")
    for name, now in (
        ("LIFE_TABLE_SOURCE", DESTATIS + href.split("?")[0].removesuffix(".xlsx") + ".html"),
        ("LIFE_TABLE_ATTRIBUTION", copyright_line(book["Impressum"])),
    ):
        if getattr(series, name) != now:
            print(f"Set series.{name} and SOURCES.md to {now}")


if __name__ == "__main__":
    delivered()
    west()
    gold()
    life_table()
